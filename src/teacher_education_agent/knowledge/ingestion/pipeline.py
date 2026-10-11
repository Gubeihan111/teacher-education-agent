"""Source-preserving DOCX, PDF, VSDX and text ingestion pipeline."""
from __future__ import annotations
import hashlib, re, zipfile
from pathlib import Path
from xml.etree import ElementTree
from sqlalchemy import delete
from teacher_education_agent.infrastructure.database.models import KnowledgeChunk

def extract_text(path: Path) -> str:
    suffix=path.suffix.lower()
    if suffix==".docx":
        from docx import Document
        doc=Document(path); parts=[p.text for p in doc.paragraphs]
        parts += ["\t".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows]
        return "\n".join(parts)
    if suffix==".pdf":
        from pypdf import PdfReader
        return "\n\f\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    if suffix==".vsdx":
        texts=[]
        with zipfile.ZipFile(path) as archive:
            for name in archive.namelist():
                if name.startswith("visio/pages/") and name.endswith(".xml"):
                    root=ElementTree.fromstring(archive.read(name))
                    texts.extend(node.text.strip() for node in root.iter() if node.text and node.text.strip())
        return "\n".join(texts)
    if suffix in {".txt",".md"}: return path.read_text(encoding="utf-8")
    raise ValueError(f"unsupported source type: {suffix}")

def chunk_text(text: str, *, size: int=900, overlap: int=120) -> list[str]:
    normalized=re.sub(r"[ \t]+"," ",text).strip(); chunks=[]; start=0
    while start < len(normalized):
        end=min(len(normalized),start+size); boundary=max(normalized.rfind("\n",start,end),normalized.rfind("。",start,end))
        if boundary > start+size//2: end=boundary+1
        chunks.append(normalized[start:end].strip())
        if end>=len(normalized): break
        start=max(start+1,end-overlap)
    return [item for item in chunks if item]

class KnowledgeIngestor:
    def __init__(self, models, sessions, *, vector_dimension: int): self.models,self.sessions,self.vector_dimension=models,sessions,vector_dimension
    async def ingest(self, path: Path, metadata: dict | None=None) -> dict:
        content=extract_text(path); chunks=chunk_text(content)
        if not chunks: raise ValueError("source contains no extractable text")
        vectors=await self.models.embed(chunks)
        if any(len(vector)!=self.vector_dimension for vector in vectors): raise ValueError("embedding dimension does not match VECTOR_DIMENSION")
        source_ref=str(path.resolve()); digest=hashlib.sha256(path.read_bytes()).hexdigest()
        async with self.sessions() as session:
            await session.execute(delete(KnowledgeChunk).where(KnowledgeChunk.source_ref==source_ref))
            for index,(chunk,vector) in enumerate(zip(chunks,vectors)):
                session.add(KnowledgeChunk(source_ref=source_ref,content=chunk,embedding=vector,metadata_json={**(metadata or {}),"chunk_index":index,"sha256":digest,"filename":path.name}))
            await session.commit()
        return {"source_ref":source_ref,"sha256":digest,"chunk_count":len(chunks)}
